# -*- coding: utf-8 -*-
"""Bangun 'BAB 1 - TANGSEL LIA.md' dan '.docx' dari satu sumber konten.

Format docx: Times New Roman 12, spasi 1.5, margin T4 L4 R3 B3 cm, Table Grid.
Markdown memakai markup yang sama (*italic*, **bold**) sehingga sinkron dengan docx.
"""
import io
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

import bab1_content as K

DOCS = os.path.dirname(os.path.abspath(__file__))
OUT_MD = os.path.join(DOCS, "BAB 1 - TANGSEL LIA.md")
OUT_DOCX = os.path.join(DOCS, "BAB 1 - TANGSEL LIA.docx")
FONT = "Times New Roman"

TOKEN = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*)", re.S)


def parse_inline(text):
    """-> list of (teks, bold, italic)."""
    parts = []
    for chunk in TOKEN.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            parts.append((chunk[2:-2], True, False))
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            parts.append((chunk[1:-1], False, True))
        else:
            parts.append((chunk, False, False))
    return parts


# --------------------------------------------------------------------------
# Markdown
# --------------------------------------------------------------------------
def build_md():
    L = []
    a = L.append
    a("# BAB I")
    a("")
    a("# PENDAHULUAN")
    a("")
    a("## 1.1 Latar Belakang")
    a("")
    for t in K.LATAR:
        a(t)
        a("")
    a("## 1.2 Rumusan Masalah")
    a("")
    a("Berdasarkan latar belakang yang telah diuraikan, maka rumusan masalah dalam penelitian ini adalah sebagai berikut:")
    a("")
    for i, t in enumerate(K.RUMUSAN, 1):
        a("%d. %s" % (i, t))
        a("")
    a("## 1.3 Tujuan Penelitian")
    a("")
    a("Berdasarkan rumusan masalah yang telah dikemukakan, maka tujuan dari pelaksanaan Tugas Akhir ini adalah:")
    a("")
    for i, t in enumerate(K.TUJUAN, 1):
        a("%d. %s" % (i, t))
        a("")
    a("## 1.4 Batasan Masalah")
    a("")
    a("Untuk menjaga agar penelitian ini tetap fokus dan terarah, batasan masalah (ruang lingkup) yang ditetapkan adalah sebagai berikut:")
    a("")
    for i, (lead, body) in enumerate(K.BATASAN, 1):
        a("%d. **%s** %s" % (i, lead, body))
        a("")
    a("## 1.5 State of the Art")
    a("")
    a("**Tabel 1.1 State of The Art**")
    a("")
    for e in K.SOTA:
        a("### %s" % e["judul"])
        a("")
        a("| | |")
        a("|---|---|")
        for k in ("Peneliti", "Lokasi", "Tahun", "Nama Jurnal", "DOI/URL",
                  "Hasil Penelitian", "Alasan Menjadi Tinjauan Penelitian"):
            a("| **%s** | %s |" % (k, e[k]))
        a("")
    a("## 1.6 Sistematika Penulisan")
    a("")
    a("Untuk memberikan gambaran yang jelas dan terstruktur mengenai penelitian ini, penyusunan Tugas Akhir dibagi menjadi lima bab dengan sistematika penulisan sebagai berikut:")
    a("")
    for bab, desc in K.SISTEMATIKA:
        a("**%s**" % bab)
        a("")
        a(desc)
        a("")
    a("## Daftar Pustaka")
    a("")
    for i, ref in enumerate(K.DAFTAR, 1):
        a("[%d] %s" % (i, ref))
        a("")
    text = "\n".join(L).rstrip() + "\n"
    with io.open(OUT_MD, "w", encoding="utf-8", newline="\n") as f:
        f.write(text)
    return text


# --------------------------------------------------------------------------
# Docx
# --------------------------------------------------------------------------
def force_font(run, size=12, bold=None, italic=None):
    run.font.name = FONT
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)


def new_doc():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(12)
    pf = style.paragraph_format
    pf.line_spacing = 1.5
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for s in doc.sections:
        s.top_margin = Cm(4)
        s.left_margin = Cm(4)
        s.right_margin = Cm(3)
        s.bottom_margin = Cm(3)
    return doc


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    force_font(run)
    r = run._element
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, sep, end):
        r.append(el)


def para(doc, text, size=12, bold=False, italic=False, align="justify",
         indent=None, space_after=6, line=1.5):
    p = doc.add_paragraph()
    p.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[align]
    pf = p.paragraph_format
    pf.line_spacing = line
    pf.space_after = Pt(space_after)
    if indent is not None:
        pf.first_line_indent = Cm(indent)
    run = p.add_run(text)
    force_font(run, size=size, bold=bold, italic=italic)
    return p


def para_rich(doc, text, align="justify", indent=None, space_after=6, line=1.5,
              size=12):
    """Teks dengan markup *italic* / **bold**."""
    p = doc.add_paragraph()
    p.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[align]
    pf = p.paragraph_format
    pf.line_spacing = line
    pf.space_after = Pt(space_after)
    if indent is not None:
        pf.first_line_indent = Cm(indent)
    for chunk, bold, italic in parse_inline(text):
        run = p.add_run(chunk)
        force_font(run, size=size, bold=bold, italic=italic)
    return p


def number_item(doc, number, text, rich=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    pf.left_indent = Cm(0.75)
    pf.first_line_indent = Cm(-0.75)
    run = p.add_run("%d.\t" % number)
    force_font(run)
    if rich:
        for chunk, bold, italic in parse_inline(text):
            r = p.add_run(chunk)
            force_font(r, bold=bold, italic=italic)
    else:
        r = p.add_run(text)
        force_font(r)
    return p


def sota_table(doc, rows):
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for label, value in rows:
        cells = table.add_row().cells
        p0 = cells[0].paragraphs[0]
        p0.paragraph_format.line_spacing = 1.15
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(label)
        force_font(r0, size=11, bold=True)
        p1 = cells[1].paragraphs[0]
        p1.paragraph_format.line_spacing = 1.15
        p1.paragraph_format.space_after = Pt(2)
        for chunk, bold, italic in parse_inline(value):
            r1 = p1.add_run(chunk)
            force_font(r1, size=11, bold=bold, italic=italic)
    table.columns[0].width = Cm(4.0)
    table.columns[1].width = Cm(10.0)
    for row in table.rows:
        row.cells[0].width = Cm(4.0)
        row.cells[1].width = Cm(10.0)


def build_docx():
    doc = new_doc()
    add_page_number(doc.sections[0].footer.paragraphs[0])

    para(doc, "BAB I", bold=True, align="center", space_after=0, line=1.5)
    para(doc, "PENDAHULUAN", bold=True, align="center", space_after=18, line=1.5)

    para(doc, "1.1 Latar Belakang", bold=True, align="left", space_after=8)
    for t in K.LATAR:
        para_rich(doc, t, indent=1.25)

    para(doc, "1.2 Rumusan Masalah", bold=True, align="left", space_after=8)
    para(doc, "Berdasarkan latar belakang yang telah diuraikan, maka rumusan masalah dalam penelitian ini adalah sebagai berikut:", space_after=8)
    for i, t in enumerate(K.RUMUSAN, 1):
        number_item(doc, i, t)

    para(doc, "1.3 Tujuan Penelitian", bold=True, align="left", space_after=8)
    para(doc, "Berdasarkan rumusan masalah yang telah dikemukakan, maka tujuan dari pelaksanaan Tugas Akhir ini adalah:", space_after=8)
    for i, t in enumerate(K.TUJUAN, 1):
        number_item(doc, i, t)

    para(doc, "1.4 Batasan Masalah", bold=True, align="left", space_after=8)
    para(doc, "Untuk menjaga agar penelitian ini tetap fokus dan terarah, batasan masalah (ruang lingkup) yang ditetapkan adalah sebagai berikut:", space_after=8)
    for i, (lead, body) in enumerate(K.BATASAN, 1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.75)
        r = p.add_run("%d.\t" % i)
        force_font(r)
        r = p.add_run(lead + " ")
        force_font(r, bold=True)
        for chunk, bold, italic in parse_inline(body):
            r = p.add_run(chunk)
            force_font(r, bold=bold, italic=italic)

    para(doc, "1.5 State of The Art", bold=True, align="left", space_after=8)
    para(doc, "Tabel 1.1 State of The Art", bold=True, align="center", space_after=10)
    for e in K.SOTA:
        para(doc, e["judul"], italic=True, align="left", space_after=6)
        rows = [(k, e[k]) for k in
                ("Peneliti", "Lokasi", "Tahun", "Nama Jurnal", "DOI/URL",
                 "Hasil Penelitian", "Alasan Menjadi Tinjauan Penelitian")]
        sota_table(doc, rows)
        para(doc, "", space_after=6)

    para(doc, "1.6 Sistematika Penulisan", bold=True, align="left", space_after=8)
    para(doc, "Untuk memberikan gambaran yang jelas dan terstruktur mengenai penelitian ini, penyusunan Tugas Akhir dibagi menjadi lima bab dengan sistematika penulisan sebagai berikut:", space_after=8)
    for bab, desc in K.SISTEMATIKA:
        para(doc, bab, bold=True, align="left", space_after=4)
        para_rich(doc, desc, indent=1.25)

    para(doc, "Daftar Pustaka", bold=True, align="center", space_after=12)
    for i, ref in enumerate(K.DAFTAR, 1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.9)
        r = p.add_run("[%d]\t" % i)
        force_font(r)
        for chunk, bold, italic in parse_inline(ref):
            r = p.add_run(chunk)
            force_font(r, bold=bold, italic=italic)

    doc.save(OUT_DOCX)


def main():
    md = build_md()
    build_docx()
    print("MD   ->", OUT_MD, os.path.getsize(OUT_MD), "bytes")
    print("DOCX ->", OUT_DOCX, os.path.getsize(OUT_DOCX), "bytes")
    print("md chars:", len(md))
    print("italic markers in md:", md.count("*") )


if __name__ == "__main__":
    main()
